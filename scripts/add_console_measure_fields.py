#!/usr/bin/env python3
"""Place the on-board measurement fields in the console templates.

The console shows the counters that the processor freezes when the program
writes to tohost (cycles_at_done / instr_at_done).  They are the on-board
counterpart of the values cocotb reports, so a photograph of the console is
comparable with the simulation numbers.

The fields are written into cells that are empty in the trimmed template, so
the layout of the rest of the screen does not move.  The script is idempotent:
it overwrites its own cells and nothing else.

Usage:  python3 scripts/add_console_measure_fields.py
"""

import openpyxl

XLSX = "scripts/RiscVScreenV3.xlsx"

# (sheet, row, first column, label) - 1-based coordinates.
# The 16 X characters are the placeholder of a 64-bit hexadecimal value.
FIELDS = [
    ("Single-cycle", 8, 2, "CyclesToPass"),
    ("Single-cycle", 9, 2, "InstrToPass"),
    ("Pipeline", 2, 45, "CyclesToPass"),
    ("Pipeline", 3, 45, "InstrToPass"),
]

VALUE_WIDTH = 16


def place(ws, row: int, col: int, label: str) -> None:
    """Write "<label>: <16 X>" starting at (row, col)."""
    text = f"{label}: "
    for i, ch in enumerate(text):
        ws.cell(row=row, column=col + i).value = ch
    for i in range(VALUE_WIDTH):
        ws.cell(row=row, column=col + len(text) + i).value = 'X'


def main() -> None:
    wb = openpyxl.load_workbook(XLSX)
    for sheet, row, col, label in FIELDS:
        place(wb[sheet], row, col, label)
    wb.save(XLSX)
    print(f"measurement fields placed in {XLSX}")


if __name__ == '__main__':
    main()
