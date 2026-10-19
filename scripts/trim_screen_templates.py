#!/usr/bin/env python3
"""Trim the RiscV console templates to the 160x45 text-mode viewport.

The 720p60 framebuffer holds exactly 160 columns x 45 rows.  Both worksheets
were drawn larger, so gen_screen_mif.py used to drop the overflow silently.

  Single-cycle : 50 rows x 161 cols -> 45 x 160
  Pipeline     : 51 rows x 161 cols -> 45 x 160

The layout is drawn with text characters, not cell borders, so rows can be
moved by copying their cell values.
"""

import openpyxl

XLSX = "scripts/RiscVScreenV3.xlsx"
COLS = 160
ROWS = 45


def row_vals(ws, r, c0, c1):
    return [ws.cell(row=r, column=c).value for c in range(c0, c1 + 1)]


def put_row(ws, r, c0, c1, vals):
    """Clear the span first, then write: stale characters must not survive."""
    for c in range(c0, c1 + 1):
        ws.cell(row=r, column=c).value = None
    for i, v in enumerate(vals):
        if v is not None:
            ws.cell(row=r, column=c0 + i).value = v


def move_rows(ws, srcs, dsts, c0, c1):
    buf = {s: row_vals(ws, s, c0, c1) for s in srcs}
    for s, d in zip(srcs, dsts):
        put_row(ws, d, c0, c1, buf[s])
    for s in srcs:
        if s not in dsts:
            put_row(ws, s, c0, c1, [None] * (c1 - c0 + 1))


def clear_rows(ws, rows, c0, c1):
    for r in rows:
        put_row(ws, r, c0, c1, [None] * (c1 - c0 + 1))


def border(c0, c1):
    vals = [None] * (c1 - c0 + 1)
    vals[0] = '|'
    vals[-1] = '|'
    for i in range(1, c1 - c0):
        vals[i] = '-'
    return vals


def trim_single_cycle(ws):
    # 1. Prose block (cols 1-48, rows 6-8) moves up one row.
    move_rows(ws, [6, 7, 8], [5, 6, 7], 1, 48)

    # 2. Register / memory tables (cols 1-48, rows 11-46) move up one row so
    #    the closing border lands on the last visible row.
    move_rows(ws, list(range(11, 47)), list(range(10, 46)), 1, 48)

    # 3. "Clock cycles" box (cols 51-83) moves up one row and widens from 8 to
    #    16 characters, so the 64-bit cycle counter fits.
    lbl = ['|']
    lbl += list("Clock cycles")
    lbl += [':', None]
    lbl += ['X'] * 16
    lbl += [None, '|']
    put_row(ws, 43, 51, 83, border(51, 83))
    put_row(ws, 44, 51, 83, lbl)
    put_row(ws, 45, 51, 83, border(51, 83))

    # 4. Row label of that box follows it.
    put_row(ws, 45, 108, 109, row_vals(ws, 46, 108, 109))

    # 5. Everything below the viewport goes away.
    clear_rows(ws, [46, 47, 48, 49, 50], 1, COLS)


def trim_pipeline(ws):
    # 1. EXECUTE box (cols 32-55): drop the separator and the blank row, so
    #    the box closes on the last visible row.
    move_rows(ws, [41, 42, 43, 45, 46, 47], [40, 41, 42, 43, 44, 45], 32, 55)

    # 2. MEM/WB box and the arrow into it (cols 56-85): drop the separator
    #    only.  The arrow row must stay on the ALURes row of the EXECUTE box,
    #    so both boxes shift by the same amount for that row.
    move_rows(ws, [41, 42, 43, 44, 45, 46], [40, 41, 42, 43, 44, 45], 56, 85)

    # 3. Counters box (cols 98-130) shrinks to three rows.
    top = row_vals(ws, 44, 98, 130)
    title = row_vals(ws, 45, 98, 130)
    value = row_vals(ws, 47, 98, 130)
    put_row(ws, 43, 98, 130, top)
    put_row(ws, 44, 98, 130, title)
    put_row(ws, 45, 98, 130, value)

    # 4. Everything below the viewport goes away.
    clear_rows(ws, [46, 47, 48, 49, 50, 51], 1, COLS)


def main():
    wb = openpyxl.load_workbook(XLSX)
    trim_single_cycle(wb["Single-cycle"])
    trim_pipeline(wb["Pipeline"])

    # The column ruler of both sheets ends with a '$' in column 161.
    for name, rows in (("Single-cycle", [2]), ("Pipeline", [3, 5])):
        for r in rows:
            wb[name].cell(row=r, column=161).value = None

    # HALTED is a single bit, so it takes a binary placeholder.
    wb["Single-cycle"].cell(row=7, column=103).value = 'B'

    wb.save(XLSX)
    print("templates trimmed to %dx%d" % (COLS, ROWS))


if __name__ == '__main__':
    main()
